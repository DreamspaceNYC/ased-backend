from fastapi import FastAPI, UploadFile, File
from fastapi.responses import StreamingResponse
import pytesseract, cv2, io, numpy as np
from fpdf import FPDF
from docx import Document
import qrcode

app = FastAPI()

@app.get("/")
def read_root():
    return {"status": "ASED backend is live"}

@app.post("/ocr")
async def ocr(file: UploadFile = File(...)):
    image = await file.read()
    npimg = cv2.imdecode(np.frombuffer(image, np.uint8), cv2.IMREAD_COLOR)
    text = pytesseract.image_to_string(npimg)
    return {"text": text}

@app.post("/export-pdf")
async def export_pdf(content: dict):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    for line in content.get("content", "").split("\n"):
        pdf.cell(200, 10, txt=line, ln=1)
    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/pdf")

@app.post("/export-docx")
async def export_docx(content: dict):
    doc = Document()
    for line in content.get("content", "").split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")

@app.post("/generate-qr")
async def generate_qr(data: dict):
    qr = qrcode.make(data.get("data", ""))
    buf = io.BytesIO()
    qr.save(buf, format="PNG")
    buf.seek(0)
    return StreamingResponse(buf, media_type="image/png")
